"""
PASSAGE Backend — FastAPI
KANTEKANT Group · B.E Company
Jéricho BOURA · ktkintel@gmail.com
v1.1.0 — MODULE TRADUCTION intégré (translate.py)
"""

from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
import tempfile, os, re, uuid
from datetime import datetime
from typing import Optional

# ── MODULE TRADUCTION (nouveau) ──
from translate import translate_designations, detect_language
from translate_router import router as translate_router

app = FastAPI(title="PASSAGE API", version="1.1.0",
              description="KANTEKANT Group — Document transformation engine")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ── Enregistrer les endpoints /translate/* ──
app.include_router(translate_router)

MARCHES_PARAMS = {
    "Maroc":       {"transport": 10.0, "douane": 2.5,  "marge": 22.0},
    "Mayotte":     {"transport": 8.0,  "douane": 5.0,  "marge": 20.0},
    "Caraïbe":     {"transport": 12.0, "douane": 5.0,  "marge": 25.0},
    "Martinique":  {"transport": 10.0, "douane": 3.0,  "marge": 22.0},
    "France":      {"transport": 6.0,  "douane": 0.0,  "marge": 18.0},
    "Guyane":      {"transport": 10.0, "douane": 3.0,  "marge": 20.0},
    "Guadeloupe":  {"transport": 10.0, "douane": 3.0,  "marge": 20.0},
}

SOCIETE = {
    "groupe": "Groupe KANTEKANT",
    "contact": "Jéricho BOURA",
    "tel": "+596 696 415 157",
    "email": "ktkintel@gmail.com",
    "adresse": "Lotissement la Trompeuse n°5, Zone de Californie, 97232 Le Lamentin, Martinique",
}

@app.get("/")
def root():
    return {
        "app": "PASSAGE",
        "version": "1.1.0",
        "groupe": "KANTEKANT",
        "status": "operational",
        "endpoints": ["/transform", "/health", "/marches", "/translate/designations",
                      "/translate/blocks", "/translate/text", "/translate/detect"]
    }

@app.get("/health")
def health():
    return {"status": "ok", "timestamp": datetime.now().isoformat()}

@app.get("/marches")
def get_marches():
    return MARCHES_PARAMS

@app.post("/transform")
async def transform_document(
    file: UploadFile = File(...),
    filiale: str = Form("B.E Energies"),
    marche: str = Form("Martinique"),
    devise_source: str = Form("USD"),
    taux_change: float = Form(0.92),
    transport_pct: Optional[float] = Form(None),
    douane_pct: Optional[float] = Form(None),
    marge_pct: Optional[float] = Form(None),
    client_nom: Optional[str] = Form(""),
    masquer_fournisseur: bool = Form(True),
):
    # Paramètres du marché
    params = dict(MARCHES_PARAMS.get(marche, MARCHES_PARAMS["Martinique"]))
    if transport_pct is not None: params["transport"] = transport_pct
    if douane_pct is not None:    params["douane"]    = douane_pct
    if marge_pct is not None:     params["marge"]     = marge_pct

    suffix = os.path.splitext(file.filename)[1].lower()
    tmp_dir = tempfile.mkdtemp()
    input_path = os.path.join(tmp_dir, f"input{suffix}")

    with open(input_path, "wb") as f:
        content = await file.read()
        f.write(content)

    try:
        if suffix == ".pdf":
            produits, titre = extraire_pdf(input_path)
        elif suffix in [".xlsx", ".xls"]:
            produits, titre = extraire_excel(input_path)
        elif suffix == ".docx":
            produits, titre = extraire_docx(input_path)
        else:
            raise HTTPException(status_code=400,
                detail=f"Format non supporté : {suffix}. Acceptés : PDF, XLSX, DOCX")
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Erreur extraction : {str(e)}")

    if not produits:
        raise HTTPException(status_code=422,
            detail="Aucun produit/prix détecté dans ce document.")

    # ── TRADUCTION GROUPÉE (un seul appel Claude pour tous les produits) ──
    # Plus efficace que traduire() appelé une fois par produit
    descs_brutes = [p.get("desc", "") for p in produits]
    descs_fr     = traduire_lot(descs_brutes)
    for p, desc_fr in zip(produits, descs_fr):
        p["desc"] = desc_fr

    for p in produits:
        p["prix_client"] = calculer_prix(
            p.get("prix_source"), devise_source, taux_change, params)

    ref = f"{filiale.replace('.','').replace(' ','-').upper()}-{marche[:3].upper()}-{datetime.now().strftime('%Y%m%d%H%M')}"
    output_path = os.path.join(tmp_dir, f"PASSAGE_{ref}.pdf")

    generer_pdf_be(
        produits=produits,
        titre_source=titre,
        filiale=filiale,
        marche=marche,
        params=params,
        devise_source=devise_source,
        taux_change=taux_change,
        client_nom=client_nom or "",
        ref=ref,
        output_path=output_path,
        societe=SOCIETE,
    )

    return FileResponse(
        output_path,
        media_type="application/pdf",
        filename=f"PASSAGE_{ref}.pdf",
        headers={"X-Produits-Count": str(len(produits)),
                 "X-Reference": ref,
                 "X-Marche": marche,
                 "X-Filiale": filiale}
    )


# ══════════════════════════════════════════════════════════════════
# TRADUCTION — intégration du module partagé
# ══════════════════════════════════════════════════════════════════

def traduire_lot(descriptions: list[str]) -> list[str]:
    """
    Traduit une liste de descriptions en un seul appel Claude.
    Remplace l'ancienne fonction traduire() appelée produit par produit.
    
    - Si toutes les descriptions sont vides → retourne les originaux
    - Si déjà en français → retourne les originaux (détection automatique)
    - Sinon → appel translate_designations() du module partagé
    """
    if not descriptions:
        return descriptions

    # Vérification rapide : si toutes vides, pas la peine d'appeler l'API
    non_vides = [d for d in descriptions if d and d.strip()]
    if not non_vides:
        return descriptions

    try:
        return translate_designations(descriptions)
    except Exception as e:
        # Fallback : si l'API échoue, utiliser l'ancienne fonction de substitution
        print(f"[PASSAGE] Fallback traduction locale (erreur API: {e})")
        return [_traduire_local(d) for d in descriptions]


def _traduire_local(t: str) -> str:
    """
    Fallback : dictionnaire de substitutions (ancienne fonction traduire).
    Utilisé uniquement si l'API Claude est indisponible.
    """
    if not t:
        return t
    subs = [
        ("Water-activated EV Power Generator for cars", "Générateur EV eau — coffre voiture"),
        ("Water-activated EV Power Generator", "Générateur EV à activation par eau"),
        ("Aluminum Plate for", "Plaques aluminium pour"),
        ("Electrolyte Powder for", "Poudre électrolyte pour"),
        ("emergency power generator", "générateur secours"),
        ("Accessories of Three Disruptions", "Système urgence 3 ruptures"),
        ("Emergency Communication and Power", "Communication & énergie d'urgence"),
        ("Salt Water", "Eau salée"), ("Flashlight", "Lampe torche"),
        ("Portable", "Portable"), ("Generator", "Générateur"),
    ]
    for en, fr in subs:
        t = t.replace(en, fr)
    return t


# ══════════════════════════════════════════════════════════════════
# EXTRACTION PDF — inchangée, sauf : traduire(desc) → desc brute
# (la traduction est maintenant faite en lot dans /transform)
# ══════════════════════════════════════════════════════════════════

def extraire_pdf(path):
    import pdfplumber
    produits, titre = [], []
    refs_vus = set()

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            if i == 0:
                for l in text.split("\n"):
                    if any(x in l for x in ["Quotation","Quote","报价","价格","Price","Devis","DEVIS"]):
                        titre = l.strip(); break

            for table in page.extract_tables():
                if not table: continue

                header_idx = None
                for hi, hrow in enumerate(table[:3]):
                    hdr = " ".join(str(c) for c in (hrow or []) if c).lower()
                    if "item" in hdr or "no." in hdr or "description" in hdr:
                        header_idx = hi
                        break
                if header_idx is None:
                    continue

                for row in table[header_idx+1:]:
                    if not row: continue
                    row = [str(c).strip() if c else "" for c in row]
                    ncols = len(row)

                    no = row[0] if row[0] else ""
                    if not no.isdigit():
                        continue

                    ref_prod = row[1].replace("\n", "-") if ncols > 1 else ""
                    ref_prod = ref_prod.strip()

                    if ref_prod and ref_prod in refs_vus:
                        continue

                    if ncols >= 5:
                        desc_raw = row[3] if row[3] else (row[2] if row[2] else ref_prod)
                    else:
                        desc_raw = row[2] if ncols > 2 and row[2] else ref_prod
                    desc = desc_raw.split("\n")[0].strip()

                    prix = None
                    for col_idx in [6, 7, 5]:
                        if ncols > col_idx and row[col_idx]:
                            prix = parse_prix_cellule(row[col_idx])
                            if prix: break

                    if not prix:
                        for cell in row:
                            if cell:
                                prix = parse_prix_cellule(cell)
                                if prix: break

                    if ref_prod or prix:
                        refs_vus.add(ref_prod)
                        produits.append({
                            "no": no,
                            "ref": ref_prod or f"REF-{no}",
                            "desc": desc,          # ← brut, traduit en lot après
                            "prix_source": prix,
                            "qte": "",
                        })

    return produits, titre or "Document fournisseur"


def parse_prix_cellule(cellule):
    if not cellule:
        return None
    text = str(cellule)
    lignes = text.split("\n")

    for ligne in lignes:
        if "sample" in ligne.lower() or "échantillon" in ligne.lower():
            p = parse_prix(ligne)
            if p and p > 0.5:
                return p

    for ligne in lignes:
        ligne_stripped = ligne.strip()
        if ligne_stripped.startswith(">=") or ligne_stripped.startswith(">"):
            continue
        p = parse_prix(ligne_stripped)
        if p and p > 0.5:
            return p

    return parse_prix(text)


def extraire_excel(path):
    import openpyxl
    produits, titre = [], ""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return produits, titre
    header_row = 0
    for i, row in enumerate(rows[:10]):
        row_str = " ".join(str(c) for c in row if c).lower()
        if any(x in row_str for x in ["price","prix","model","item","产品"]):
            header_row = i; break
    for row in rows[header_row+1:]:
        if not any(row): continue
        cells = [str(c).strip() if c is not None else "" for c in row]
        prix = None
        for cell in cells:
            p = parse_prix(cell)
            if p and p > 0: prix = p; break
        desc = next((c for c in cells if len(c)>3 and
                     not c.replace(".","").replace(",","").isdigit()), "")
        if desc or prix:
            produits.append({
                "no": str(len(produits)+1),
                "ref": cells[0][:20] if cells[0] else f"REF-{len(produits)+1}",
                "desc": desc,          # ← brut, traduit en lot après
                "prix_source": prix,
                "qte": "",
            })
    return produits, titre or "Document fournisseur"


def extraire_docx(path):
    from docx import Document
    produits, titre = [], ""
    doc = Document(path)
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            if i == 0: continue
            cells = [c.text.strip() for c in row.cells]
            if not any(cells): continue
            prix = None
            for cell in cells:
                p = parse_prix(cell)
                if p and p > 0: prix = p; break
            desc = next((c for c in cells if len(c)>3 and
                        not c.replace(".","").replace(",","").isdigit()), "")
            if desc or prix:
                produits.append({
                    "no": str(i),
                    "ref": cells[0][:20] if cells else f"REF-{i}",
                    "desc": desc,          # ← brut, traduit en lot après
                    "prix_source": prix,
                    "qte": "",
                })
    return produits, titre or "Document fournisseur"


# ══════════════════════════════════════════════════════════════════
# UTILITAIRES — inchangés
# ══════════════════════════════════════════════════════════════════

def parse_prix(s):
    if not s: return None
    m = re.search(r"[\$¥￥]?\s*([\d,]+\.?\d*)", str(s))
    if m:
        try:
            v = float(m.group(1).replace(",",""))
            if v > 0.5: return v
        except: pass
    return None

def parse_prix_bulk(cellule):
    if not cellule: return None
    for l in str(cellule).split("\n"):
        l = l.strip()
        if l.startswith(">="): continue
        clean = l.replace("$","").replace("/set","").replace(",","").strip()
        try:
            v = float(clean)
            if v > 5: return v
        except: pass
        m = re.search(r"([\d,]+\.?\d*)/set", l)
        if m:
            try:
                v = float(m.group(1).replace(",",""))
                if v > 5: return v
            except: pass
    return None

def calculer_prix(prix_source, devise, taux, params):
    if not prix_source: return None
    eur = prix_source * taux
    eur_tr = eur * (1 + params["transport"]/100)
    eur_dou = eur_tr * (1 + params["douane"]/100)
    return round(eur_dou * (1 + params["marge"]/100), 2)


# ══════════════════════════════════════════════════════════════════
# GÉNÉRATION PDF — inchangée
# ══════════════════════════════════════════════════════════════════

def generer_pdf_be(produits, titre_source, filiale, marche, params,
                    devise_source, taux_change, client_nom, ref, output_path, societe):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                     Table, TableStyle, HRFlowable)
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT

    ROUGE  = colors.HexColor("#CC0000")
    NOIR   = colors.HexColor("#111111")
    GRIS   = colors.HexColor("#666666")
    GRIS_L = colors.HexColor("#f7f7f7")
    GRIS_B = colors.HexColor("#dddddd")
    VERT   = colors.HexColor("#1a5c3a")
    BLEU   = colors.HexColor("#1a3a5c")

    doc = SimpleDocTemplate(output_path, pagesize=A4,
        rightMargin=1.8*cm, leftMargin=1.8*cm,
        topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    S = lambda name, **kw: ParagraphStyle(name, parent=styles["Normal"], **kw)
    story = []

    hd = [[
        Paragraph(filiale, S("n", fontSize=20, fontName="Helvetica-Bold",
                             textColor=NOIR, spaceAfter=2)),
        Paragraph(f"Réf. {ref}", S("r", fontSize=8, textColor=GRIS,
                                    alignment=TA_RIGHT)),
    ]]
    ht = Table(hd, colWidths=[11*cm, 6.7*cm])
    ht.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LEFTPADDING",(0,0),(-1,-1),0),
        ("RIGHTPADDING",(0,0),(-1,-1),0),
        ("BOTTOMPADDING",(0,0),(-1,-1),0),
    ]))
    story.append(ht)
    story.append(Paragraph("Sourcing · Distribution · Solutions",
        S("t", fontSize=9, textColor=GRIS, spaceAfter=4)))
    story.append(HRFlowable(width="100%", thickness=2.5, color=ROUGE, spaceAfter=3))
    story.append(HRFlowable(width="100%", thickness=0.4, color=GRIS_B, spaceAfter=10))

    story.append(Paragraph(f"Offre commerciale — {marche}",
        S("oc", fontSize=13, fontName="Helvetica-Bold", textColor=ROUGE,
          alignment=TA_CENTER, spaceAfter=3)))
    if client_nom:
        story.append(Paragraph(f"Préparé pour : {client_nom}",
            S("cl", fontSize=10, textColor=BLEU, alignment=TA_CENTER, spaceAfter=3)))
    story.append(Paragraph(
        f"Tarifs incluant transport ({params['transport']}%), "
        f"droits de douane ({params['douane']}%) et frais de service ({params['marge']}%). "
        f"Taux appliqué : 1 {devise_source} = {taux_change} EUR. "
        f"Date : {datetime.now().strftime('%d/%m/%Y')}",
        S("note", fontSize=8.5, textColor=GRIS, spaceAfter=14)))

    sym = "€"
    data = [["N°", "Référence", "Description", f"Prix client ({sym})"]]
    for p in produits:
        pc = p.get("prix_client")
        pc_str = f"{sym} {pc:,.2f}".replace(",", " ") if pc else "Sur devis"
        ref_c  = p["ref"][:22] if len(p["ref"]) > 22 else p["ref"]
        desc_c = (p["desc"][:60]+"…") if len(p["desc"]) > 60 else p["desc"]
        data.append([p["no"], ref_c, desc_c, pc_str])

    t = Table(data, colWidths=[1.1*cm, 4.4*cm, 8.5*cm, 3.7*cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0), BLEU),
        ("TEXTCOLOR",(0,0),(-1,0), colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,0),9),
        ("ALIGN",(0,0),(-1,0),"CENTER"),
        ("TOPPADDING",(0,0),(-1,0),7),("BOTTOMPADDING",(0,0),(-1,0),7),
        ("FONTNAME",(0,1),(-1,-1),"Helvetica"),("FONTSIZE",(0,1),(-1,-1),8),
        ("ALIGN",(0,1),(1,-1),"CENTER"),("ALIGN",(3,1),(3,-1),"RIGHT"),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("TOPPADDING",(0,1),(-1,-1),6),("BOTTOMPADDING",(0,1),(-1,-1),6),
        ("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),
        *[("BACKGROUND",(0,i),(-1,i), GRIS_L) for i in range(2,len(data),2)],
        ("GRID",(0,1),(-1,-1),0.3, GRIS_B),
        ("LINEBELOW",(0,0),(-1,0),1.5, ROUGE),
        ("FONTNAME",(3,1),(3,-1),"Helvetica-Bold"),
        ("TEXTCOLOR",(3,1),(3,-1), VERT),
    ]))
    story.append(t)
    story.append(Spacer(1,14))

    story.append(HRFlowable(width="100%",thickness=0.5,color=GRIS_B,spaceAfter=8))
    conds = [
        ("Paiement", "50% acompte à la commande · Solde avant expédition"),
        ("Délais", "15–45 jours selon produit et disponibilité"),
        ("Validité", "30 jours à compter de la date d'émission"),
        ("Sourcing", "Sélection, négociation et suivi qualité assurés par " + filiale),
    ]
    cd = Table([[Paragraph(k, S(f"ck{i}", fontSize=8, fontName="Helvetica-Bold",
                                textColor=ROUGE)),
                 Paragraph(v, S(f"cv{i}", fontSize=8, textColor=GRIS))]
                for i,(k,v) in enumerate(conds)],
               colWidths=[3.2*cm, 14.5*cm])
    cd.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING",(0,0),(-1,-1),0),
        ("LINEBELOW",(0,0),(-1,-2),0.2,GRIS_B),
    ]))
    story.append(cd)
    story.append(Spacer(1,14))
    story.append(HRFlowable(width="100%",thickness=2,color=ROUGE,spaceAfter=5))
    story.append(Paragraph(
        f"{societe['groupe']} · {filiale} · {societe['email']} · {societe['tel']}",
        S("ft", fontSize=7, textColor=GRIS_B, alignment=TA_CENTER)))

    doc.build(story)
